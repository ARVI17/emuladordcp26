import csv
import io
import json
import re
import threading
import unicodedata
import warnings
import xml.etree.ElementTree as ET
import zipfile
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.models.entities import Player
from app.schemas.api import PlayerCreate

# Export the complete editable player state, including edition and variant
# metadata. Lifecycle IDs/deletion flags belong to database backups.
FIELDS = list(PlayerCreate.model_fields)
REQUIRED = {"name", "rating", "position", "league", "club", "nationality"}
BOOLEAN_FIELDS = {"rare", "special", "tradable", "protected", "starter", "duplicate_untradeable", "is_loan", "is_evolution", "is_holographic", "data_reviewed"}
JSON_FIELDS = {"alternate_positions", "attributes", "playstyles", "playstyles_plus", "roles"}
MAX_IMPORT_ROWS = 10000
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_PIXELS = 12_000_000
MAX_OCR_PIXELS = 6_000_000
OCR_TIMEOUT_SECONDS = 15
_OCR_GATE = threading.BoundedSemaphore(1)


def _csv_safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _export_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def export_players(db: Session, excel: bool = False) -> tuple[bytes, str]:
    from app.services.editions import current_version
    rows = [
        {field: _export_value(getattr(player, field, None)) for field in FIELDS}
        for player in db.scalars(select(Player).where(Player.deleted.is_(False), Player.game_version == current_version(db))).all()
    ]
    if excel:
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inventario"
        sheet.append(FIELDS)
        for row in rows:
            sheet.append([row[field] for field in FIELDS])
            # Explicit text cells prevent formula interpretation without changing
            # the underlying text, making XLSX exports lossless on reimport.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=FIELDS)
    writer.writeheader()
    writer.writerows({key: _csv_safe(value) for key, value in row.items()} for row in rows)
    return ("\ufeff" + output.getvalue()).encode("utf-8"), "text/csv; charset=utf-8"


def _headers(values) -> list[str]:
    headers = [str(value or "").strip().lower() for value in values]
    if not headers or not any(headers):
        raise ValueError("El archivo no contiene encabezados")
    if any(not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError("Los encabezados no pueden estar vacíos ni repetidos")
    missing = REQUIRED - set(headers)
    if missing:
        raise ValueError("Faltan columnas: " + ", ".join(sorted(missing)))
    unknown = set(headers) - set(FIELDS)
    if unknown:
        raise ValueError("Columnas desconocidas: " + ", ".join(sorted(unknown)))
    return headers


def _strict_bool(value, field: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "sí", "si", "yes"}:
        return True
    if normalized in {"0", "false", "no"}:
        return False
    raise ValueError(f"{field}: usa true/false, sí/no o 1/0")


def normalize_player_record(raw: dict, *, row_number: int | None = None) -> dict:
    """Apply the same player validation as manual creation, with row context.

    Omitted editions remain omitted so the route can assign the active edition
    while binding the reviewed file to its confirmation token.
    """
    try:
        cleaned = {}
        explicit_edition = bool(str(raw.get("game_version") or "").strip())
        for field, value in raw.items():
            if field not in FIELDS:
                continue
            if isinstance(value, str):
                value = value.strip()
                if value.startswith("'") and value[1:].lstrip().startswith(("=", "+", "-", "@")):
                    value = value[1:]
            if value in (None, ""):
                if field in REQUIRED:
                    raise ValueError(f"{field}: el campo es obligatorio")
                continue
            if field in BOOLEAN_FIELDS:
                value = _strict_bool(value, field)
            if field in JSON_FIELDS and isinstance(value, str):
                try:
                    value = json.loads(value)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f"{field}: debe contener JSON válido") from exc
            cleaned[field] = value
        result = PlayerCreate.model_validate(cleaned).model_dump()
        # Also reject whitespace-only data when running against older schemas.
        for field in REQUIRED - {"rating"}:
            if not str(result[field]).strip():
                raise ValueError(f"{field}: el campo es obligatorio")
        if not explicit_edition:
            result.pop("game_version", None)
        return result
    except ValidationError as exc:
        error = exc.errors()[0]
        field = ".".join(map(str, error["loc"]))
        message = f"{field}: {error['msg']}"
    except (ValueError, TypeError) as exc:
        message = str(exc)
    prefix = f"Fila {row_number}: " if row_number is not None else ""
    raise ValueError(prefix + message)


def parse_import(content: bytes, filename: str) -> list[dict]:
    extension = filename.rsplit(".", 1)[-1].lower()
    workbook = None
    try:
        if extension == "xlsx":
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                        raise ValueError("El XLSX descomprimido supera el límite permitido")
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            except (zipfile.BadZipFile, KeyError, OSError, InvalidFileException, ET.ParseError) as exc:
                raise ValueError("Archivo XLSX no válido") from exc
            rows = workbook.active.iter_rows()
            first = next(rows, None)
            if first is None:
                raise ValueError("Archivo XLSX vacío")
            headers = _headers([cell.value for cell in first])
            def values():
                for cells in rows:
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("El XLSX contiene fórmulas; importa una copia con valores")
                    yield [cell.value for cell in cells]
            raw_rows = values()
        elif extension == "csv":
            try:
                reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            except UnicodeDecodeError as exc:
                raise ValueError("El CSV debe estar codificado en UTF-8") from exc
            headers = _headers(next(reader, []))
            raw_rows = reader
        else:
            raise ValueError("Formato no permitido. Usa CSV o XLSX")
        records = []
        errors = []
        for row_number, values in enumerate(raw_rows, 2):
            if not any(value not in (None, "") for value in values):
                continue
            if len(values) != len(headers):
                raise ValueError(f"Fila {row_number}: la cantidad de columnas no coincide con el encabezado")
            if row_number > MAX_IMPORT_ROWS + 1:
                raise ValueError(f"El archivo supera {MAX_IMPORT_ROWS} jugadores")
            try:
                records.append(normalize_player_record(dict(zip(headers, values)), row_number=row_number))
            except ValueError as exc:
                errors.append(str(exc))
        if errors:
            raise ValueError('; '.join(errors[:100]) + (f'; {len(errors)} filas inválidas en total' if len(errors) > 100 else ''))
        if not records:
            raise ValueError("El archivo no contiene jugadores")
        return records
    except (zipfile.BadZipFile, KeyError, OSError, ET.ParseError) as exc:
        raise ValueError("Archivo de importación no válido") from exc
    finally:
        if workbook is not None:
            workbook.close()


def validate_image(content: bytes) -> tuple[int, int]:
    from PIL import Image, UnidentifiedImageError
    if not content or len(content) > MAX_IMAGE_BYTES:
        raise ValueError("La imagen supera el límite de 10 MB o está vacía")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(content)) as image:
                width, height = image.size
                if image.format not in {"PNG", "JPEG", "WEBP"} or width * height > MAX_IMAGE_PIXELS:
                    raise ValueError("Formato o dimensiones de imagen no permitidos")
                if width <= 0 or height <= 0:
                    raise ValueError("Dimensiones de imagen no válidas")
                image.verify()
        return width, height
    except (UnidentifiedImageError, OSError, SyntaxError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ValueError("Imagen no válida o demasiado grande") from exc


def ocr_image(content: bytes) -> dict:
    import cv2
    import numpy as np
    import pytesseract
    from PIL import Image
    width, height = validate_image(content)
    if not _OCR_GATE.acquire(blocking=False):
        raise ValueError("Hay otra lectura OCR en curso; vuelve a intentarlo")
    try:
        image = cv2.imdecode(np.frombuffer(content, np.uint8), cv2.IMREAD_GRAYSCALE)
        if image is None:
            raise ValueError("Imagen no válida")
        scale = min(1.6, (MAX_OCR_PIXELS / (width * height)) ** 0.5)
        if scale != 1:
            image = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA if scale < 1 else cv2.INTER_CUBIC)
        processed = cv2.threshold(image, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
        try:
            text = pytesseract.image_to_string(Image.fromarray(processed), lang="spa+eng", timeout=OCR_TIMEOUT_SECONDS)
        except (RuntimeError, pytesseract.TesseractError, pytesseract.TesseractNotFoundError) as exc:
            raise ValueError("No se pudo completar OCR dentro del tiempo permitido") from exc
    finally:
        _OCR_GATE.release()
    ratings = [int(value) for value in re.findall(r"\b([4-9][0-9])\b", text)]
    positions = re.findall(r"\b(POR|LD|DFC|LI|CAD|CAI|MCD|MC|MCO|MD|MI|ED|EI|SD|DC)\b", text.upper())
    lines = [line.strip() for line in text.splitlines() if len(line.strip()) > 2]
    return {"raw_text": text[:5000], "candidates": [{"name": lines[0] if lines else "", "rating": ratings[0] if ratings else None, "position": positions[0] if positions else "", "confidence": {"name": 0.45 if lines else 0.0, "rating": 0.7 if ratings else 0.0, "position": 0.75 if positions else 0.0}}], "requires_confirmation": True}

OCR_UI_WORDS = {
    "club", "plantilla", "buscar", "filtros", "ordenar", "calidad", "posicion", "liga",
    "nacionalidad", "precio", "quimica", "transferible", "intransferible", "volver", "menu",
}
PLAYER_POSITIONS = {"POR", "LD", "DFC", "LI", "CAD", "CAI", "MCD", "MC", "MCO", "MD", "MI", "ED", "EI", "SD", "DC"}
PLAYER_STAT = re.compile(r"^(?:SAL|PAR|SAQ|REF|VEL|POS|RIT|TIR|PAS|REG|DEF|F[IÍ]S)\s*\d{1,3}$", re.IGNORECASE)

def _plain(value: str) -> str:
    return unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii").lower()

def _candidate(name: str, rating: int, position: str, page: int, confidence: float) -> dict:
    return {
        "name": name[:100], "rating": rating, "position": position, "source_page": page,
        "confidence": {"name": confidence, "rating": confidence, "position": confidence},
    }

def extract_accessibility_inventory_candidates(content: bytes | str, page: int = 1) -> list[dict]:
    """Extract Companion player rows from a bounded, temporary UIAutomator tree."""
    if isinstance(content, bytes):
        if len(content) > 2_000_000: return []
        content = content.decode("utf-8", "replace")
    if len(content) > 2_000_000 or "<!DOCTYPE" in content.upper():
        return []
    try:
        root = ET.fromstring(content)
    except (ET.ParseError, ValueError):
        return []
    screen_text = [node.attrib.get("text", "").strip() for node in root.iter("node")]
    if not any(_plain(value) == "futb. de mi club" for value in screen_text):
        return []
    candidates: list[dict] = []
    for list_view in root.iter("node"):
        if list_view.attrib.get("class") != "android.widget.ListView": continue
        for row in list(list_view):
            values = [node.attrib.get("text", "").strip() for node in row.iter("node") if node.attrib.get("text", "").strip()]
            rating_index = next((index for index, value in enumerate(values) if value.isdigit() and 40 <= int(value) <= 99), None)
            if rating_index is None: continue
            position_index = next((index for index in range(rating_index + 1, len(values)) if values[index].upper() in PLAYER_POSITIONS), None)
            if position_index is None: continue
            name = ""
            for value in values[position_index + 1:]:
                plain = _plain(value)
                if not plain or value.isdigit() or value.upper() in PLAYER_POSITIONS or PLAYER_STAT.fullmatch(value): continue
                if any(word == plain or plain.startswith(f"{word} ") for word in OCR_UI_WORDS): continue
                if re.search(r"[A-Za-zÀ-ÿ]", value):
                    name = value.strip(" -|:")
                    break
            if name:
                candidates.append(_candidate(name, int(values[rating_index]), values[position_index].upper(), page, 0.99))
            if len(candidates) >= 30: return candidates
        if candidates: break
    return candidates

def find_accessibility_inventory_action(content: bytes | str, label: str) -> tuple[int, int, int, int] | None:
    if isinstance(content, bytes):
        if len(content) > 2_000_000: return None
        content = content.decode("utf-8", "replace")
    if len(content) > 2_000_000 or "<!DOCTYPE" in content.upper():
        return None
    try:
        root = ET.fromstring(content)
    except (ET.ParseError, ValueError):
        return None
    if not any(_plain(node.attrib.get("text", "").strip()) == "futb. de mi club" for node in root.iter("node")):
        return None
    expected = _plain(label)
    matches = []
    for node in root.iter("node"):
        text = _plain(node.attrib.get("text", "").strip())
        if expected != text or node.attrib.get("class") != "android.widget.Button": continue
        if node.attrib.get("clickable") != "true" or node.attrib.get("enabled") != "true": continue
        if node.attrib.get("visible-to-user") == "false": continue
        match = re.fullmatch(r"\[(\d+),(\d+)\]\[(\d+),(\d+)\]", node.attrib.get("bounds", ""))
        if match: matches.append(tuple(int(value) for value in match.groups()))
    return matches[0] if len(matches) == 1 else None

def extract_inventory_candidates(text: str, page: int = 1) -> list[dict]:
    lines = [re.sub(r"\s+", " ", line).strip(" -|:") for line in text.splitlines() if len(line.strip()) > 1]
    candidates: list[dict] = []
    inline = re.compile(r"^(?P<rating>[4-9][0-9])\s+(?P<position>POR|LD|DFC|LI|CAD|CAI|MCD|MC|MCO|MD|MI|ED|EI|SD|DC)\s+(?P<name>.+)$", re.IGNORECASE)
    index = 0
    while index < len(lines) and len(candidates) < 30:
        match = inline.fullmatch(lines[index])
        if match:
            name = match.group("name").strip()
            if re.search(r"[A-Za-zÀ-ÿ]", name):
                candidates.append(_candidate(name, int(match.group("rating")), match.group("position").upper(), page, 0.72))
            index += 1
            continue
        if index + 2 < len(lines) and re.fullmatch(r"[4-9][0-9]", lines[index]) and lines[index + 1].upper() in PLAYER_POSITIONS:
            name = lines[index + 2].strip()
            plain = _plain(name)
            if re.search(r"[A-Za-zÀ-ÿ]", name) and not PLAYER_STAT.fullmatch(name) and not any(word == plain or plain.startswith(f"{word} ") for word in OCR_UI_WORDS):
                candidates.append(_candidate(name, int(lines[index]), lines[index + 1].upper(), page, 0.72))
                index += 3
                continue
        index += 1
    return candidates

def ocr_inventory(content: bytes, page: int = 1, hierarchy: bytes | str | None = None) -> dict:
    structured = extract_accessibility_inventory_candidates(hierarchy, page) if hierarchy else []
    if structured:
        summary = "\n".join(f"{item['rating']} {item['position']} {item['name']}" for item in structured)
        return {"raw_text": summary, "candidates": structured, "mode": "inventory_review", "requires_confirmation": True, "extraction_source": "accessibility", "ocr_used": False}
    result = ocr_image(content)
    result["candidates"] = extract_inventory_candidates(result["raw_text"], page)
    result["mode"] = "inventory_review"
    result["extraction_source"] = "ocr"
    result["ocr_used"] = True
    return result

def review_dcp_requirements(text: str) -> dict:
    """Conservatively parse explicit comparators; never turn maxima into minima."""
    plain = _plain(text)
    requirements: dict[str, int] = {}
    warnings: list[str] = []
    labels = {
        "squad_size": r"numero de jugadores|jugadores en la plantilla|tamano de plantilla",
        "rating": r"valoracion(?: de (?:la )?plantilla)?|media(?: de (?:la )?plantilla)?",
        "chemistry": r"quimica(?: total)?",
        "rare": r"jugadores raros|raros",
        "special": r"jugadores especiales|especiales",
        "leagues": r"ligas|campeonatos",
        "clubs": r"clubes",
        "nations": r"paises|naciones|nacionalidades",
    }
    comparator = r"(?P<cmp>min(?:imo)?\.?|max(?:imo)?\.?|exactamente|exactos?|igual a|al menos|como maximo|como minimo|>=|<=|=)"
    for key, label in labels.items():
        pattern = rf"\b(?:{label})\b\s*[:\-]?\s*(?:{comparator}\s*[:\.]?\s*)?(?P<value>\d{{1,3}})\b"
        blocked = False
        for match in re.finditer(pattern, plain):
            value = int(match.group("value"))
            comparison = (match.group("cmp") or "").rstrip(".")
            if key == "squad_size":
                if comparison and comparison not in {"exactamente", "exacto", "exactos", "igual a", "="}:
                    warnings.append("Revisa el tamaño de plantilla: el comparador no describe una cantidad exacta")
                    blocked = True
                elif not 1 <= value <= 23:
                    warnings.append("Tamaño de plantilla fuera del rango permitido")
                    blocked = True
                elif key in requirements and requirements[key] != value:
                    warnings.append("Se detectaron tamaños de plantilla contradictorios")
                    blocked = True
                elif not blocked:
                    requirements[key] = value
                if blocked:
                    requirements.pop(key, None)
                continue
            if not comparison:
                warnings.append(f"Revisa el comparador de {key}; no se asumió un mínimo")
                continue
            if comparison in {"max", "maximo", "como maximo", "<="}:
                prefixes = ("max",)
            elif comparison in {"exactamente", "exacto", "exactos", "igual a", "="}:
                prefixes = ("min", "max")
            else:
                prefixes = ("min",)
            if "max" in prefixes and key not in {"leagues", "clubs", "nations"}:
                warnings.append(f"El requisito max_{key} necesita revisión; el optimizador no admite ese límite")
                blocked = True
                requirements.pop(f"min_{key}", None)
                continue
            limit = 99 if key == "rating" else 69 if key == "chemistry" else 23
            if not 0 <= value <= limit:
                warnings.append(f"Valor fuera de rango para {key}")
                continue
            if blocked:
                continue
            for prefix in prefixes:
                field = f"{prefix}_{key}"
                previous = requirements.get(field, value)
                requirements[field] = max(previous, value) if prefix == "min" else min(previous, value)
        minimum, maximum = requirements.get(f"min_{key}"), requirements.get(f"max_{key}")
        if minimum is not None and maximum is not None and minimum > maximum:
            requirements.pop(f"min_{key}")
            requirements.pop(f"max_{key}")
            warnings.append(f"Se detectaron límites contradictorios para {key}; revisa ambos")
    if "squad_size" not in requirements and requirements:
        requirements["squad_size"] = 11
        warnings.append("Confirma el tamaño de plantilla propuesto de 11 jugadores")
    return {"requirements": requirements, "warnings": warnings, "requires_human_review": True}


def parse_dcp_requirements(text: str) -> dict:
    return review_dcp_requirements(text)["requirements"]
