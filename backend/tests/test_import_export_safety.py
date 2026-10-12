import io

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.db.session import SessionLocal
from app.models.entities import Player
from app.services.editions import current_version
from app.services.import_export import (
    OCR_TIMEOUT_SECONDS,
    export_players,
    normalize_player_record,
    ocr_image,
    parse_import,
    validate_image,
)


BASE = {"name": "Jugador", "rating": "86", "position": "MC", "league": "Liga", "club": "Club", "nationality": "Colombia"}


@pytest.mark.parametrize(("field", "value"), [("price", "-500"), ("chemistry", "99")])
def test_import_rejects_ranges_with_row_context(field, value):
    data = f"name,rating,position,league,club,nationality,{field}\nJugador,86,MC,Liga,Club,Colombia,{value}\n".encode()
    with pytest.raises(ValueError, match=f"Fila 2: {field}"):
        parse_import(data, "club.csv")


def test_import_rejects_blank_required_and_unknown_boolean():
    with pytest.raises(ValueError, match="name"):
        normalize_player_record({**BASE, "name": "   "})
    with pytest.raises(ValueError, match="protected"):
        normalize_player_record({**BASE, "protected": "maybe"})
    valid = normalize_player_record({**BASE, "protected": " sí ", "rare": " FALSE "})
    assert valid["protected"] is True
    assert valid["rare"] is False
    assert "game_version" not in valid
    assert "game_version" not in normalize_player_record({**BASE, "game_version": "   "})


def test_import_header_validation_and_explicit_edition():
    with pytest.raises(ValueError, match="repetidos"):
        parse_import(b"name,name,rating,position,league,club,nationality\nA,B,86,MC,Liga,Club,Colombia", "club.csv")
    with pytest.raises(ValueError, match="game_version"):
        normalize_player_record({**BASE, "game_version": "FC26"})


@pytest.mark.parametrize("excel", [False, True])
def test_export_roundtrip_all_metadata_and_formula_text(client, excel):
    metadata = {
        "starter": True, "squad_role": "reserva", "usefulness": 67,
        "card_id": "variant-123", "base_card_id": "base-12",
        "alternate_positions": ["MCO", "MCD"],
        "attributes": {"passing": 90, "defending": 76},
        "playstyles": ["Pase"], "playstyles_plus": ["Pase+"],
        "roles": {"MC:Organizador": "++"}, "skill_moves": 4, "weak_foot": 5,
        "is_holographic": True, "gallery_score": 27, "data_source": "manual", "data_reviewed": True,
    }
    with SessionLocal() as db:
        player = Player(**{**BASE, "name": "=1+1", "rating": 86, "game_version": current_version(db), **metadata})
        db.add(player)
        db.commit()
        content, mime = export_players(db, excel)
    if excel:
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        sheet = workbook.active
        name_column = [cell.value for cell in sheet[1]].index("name")
        formula_name = next(row[name_column] for row in sheet.iter_rows(min_row=2) if row[name_column].value == "=1+1")
        assert formula_name.data_type == "s"
        workbook.close()
        filename = "club.xlsx"
    else:
        assert "'=1+1" in content.decode("utf-8-sig")
        filename = "club.csv"
    restored = next(row for row in parse_import(content, filename) if row["name"] == "=1+1")
    for key, value in metadata.items():
        assert restored[key] == value


def test_image_pixel_limit_precedes_ocr_decode():
    content = io.BytesIO()
    Image.new("L", (4000, 4000), color=255).save(content, "PNG")
    with pytest.raises(ValueError, match="dimensiones"):
        validate_image(content.getvalue())


def test_ocr_timeout_is_bounded_and_releases_worker(monkeypatch):
    import pytesseract
    content = io.BytesIO()
    Image.new("RGB", (100, 100), color="white").save(content, "PNG")
    calls = []

    def timed_out(*args, **kwargs):
        calls.append(kwargs)
        raise RuntimeError("Tesseract process timeout")

    monkeypatch.setattr(pytesseract, "image_to_string", timed_out)
    with pytest.raises(ValueError, match="tiempo"):
        ocr_image(content.getvalue())
    assert calls[0]["timeout"] == OCR_TIMEOUT_SECONDS
    monkeypatch.setattr(pytesseract, "image_to_string", lambda *args, **kwargs: "Jugador 86 MC")
    assert "Jugador" in ocr_image(content.getvalue())["raw_text"]
